#!/usr/bin/env python3
"""Isolated DOCX normalization worker. Reads/writes files only when Rust supplies paths."""
from __future__ import annotations
import argparse, html, json, re, sys, textwrap, uuid, zipfile
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from bs4 import BeautifulSoup, NavigableString, Tag

SCENE_BREAK = re.compile(r"^\s*(?:\*\s*\*\s*\*|#{1,3}|~\s*~\s*~)\s*$")

def run_html(paragraph) -> str:
    pieces=[]
    for run in paragraph.runs:
        value=html.escape(run.text).replace("\n","<br>")
        if not value: continue
        if run.bold: value=f"<strong>{value}</strong>"
        if run.italic: value=f"<em>{value}</em>"
        if run.underline: value=f"<u>{value}</u>"
        pieces.append(value)
    return "".join(pieces)

def inspect_docx(path: Path) -> dict:
    document=Document(path)
    chapters=[]; current_chapter=None; current_scene=None; warnings=[]
    def chapter(title="Imported Manuscript"):
        nonlocal current_chapter,current_scene
        current_chapter={"title":title or "Untitled Chapter","scenes":[]};chapters.append(current_chapter);current_scene=None
    def scene(title=None):
        nonlocal current_scene
        if current_chapter is None: chapter()
        current_scene={"title":title or f"Scene {len(current_chapter['scenes'])+1}","blocks":[]};current_chapter["scenes"].append(current_scene)
    for paragraph in document.paragraphs:
        text=paragraph.text.strip(); style=(paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading 1") or style in {"title","chapter title"}: chapter(text); continue
        if style.startswith("heading 2") or style in {"subtitle","scene heading"}: scene(text); continue
        if SCENE_BREAK.match(text): scene(); continue
        if current_scene is None: scene()
        if not text:
            current_scene["blocks"].append("<p></p>");continue
        tag="blockquote" if "quote" in style else "p"
        current_scene["blocks"].append(f"<{tag}>{run_html(paragraph) or html.escape(paragraph.text)}</{tag}>")
    if not chapters: chapter();scene();warnings.append("The document contained no paragraphs; an empty opening scene was created.")
    for ch in chapters:
        if not ch["scenes"]: current_chapter=ch;scene()
        for sc in ch["scenes"]:
            sc["html"]="".join(sc.pop("blocks")) or "<p></p>"
            sc["wordCount"]=len(re.findall(r"\b[\w’'-]+\b",BeautifulSoup(sc["html"],"html.parser").get_text(" ")))
    return {"formatVersion":1,"sourcePath":str(path),"chapters":chapters,"warnings":warnings,"counts":{"chapters":len(chapters),"scenes":sum(len(ch["scenes"]) for ch in chapters),"words":sum(sc["wordCount"] for ch in chapters for sc in ch["scenes"])}}

def finalize(chapters, path, warnings=None):
    warnings=warnings or []
    if not chapters: chapters=[{"title":"Imported Manuscript","scenes":[{"title":"Scene 1","html":"<p></p>"}]}]
    for chapter in chapters:
        if not chapter.get("scenes"):chapter["scenes"]=[{"title":"Scene 1","html":"<p></p>"}]
        for scene in chapter["scenes"]:
            scene["html"]=scene.get("html") or "<p></p>";scene["wordCount"]=len(re.findall(r"\b[\w’'-]+\b",BeautifulSoup(scene["html"],"html.parser").get_text(" ")))
    return {"formatVersion":1,"sourcePath":str(path),"chapters":chapters,"warnings":warnings,"counts":{"chapters":len(chapters),"scenes":sum(len(ch["scenes"]) for ch in chapters),"words":sum(scene["wordCount"] for chapter in chapters for scene in chapter["scenes"])}}

def markdown_inline(value):
    value=html.escape(value);value=re.sub(r"\*\*([^*]+)\*\*",r"<strong>\1</strong>",value);value=re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)",r"<em>\1</em>",value);value=re.sub(r"__([^_]+)__",r"<strong>\1</strong>",value);value=re.sub(r"(?<!_)_([^_]+)_(?!_)",r"<em>\1</em>",value);return value

def inspect_lines(path, markdown=False):
    lines=path.read_text(encoding="utf-8-sig").splitlines();chapters=[];chapter=None;scene=None;paragraph=[]
    def ensure_chapter(title="Imported Manuscript"):
        nonlocal chapter
        if chapter is None:chapter={"title":title,"scenes":[]};chapters.append(chapter)
    def ensure_scene(title=None):
        nonlocal scene
        ensure_chapter();scene={"title":title or f"Scene {len(chapter['scenes'])+1}","blocks":[]};chapter["scenes"].append(scene)
    def flush():
        nonlocal paragraph
        if not paragraph:return
        if scene is None:ensure_scene()
        raw=" ".join(part.strip() for part in paragraph);tag="blockquote" if raw.startswith(">") else "p";raw=raw.lstrip("> ") if tag=="blockquote" else raw;body=markdown_inline(raw) if markdown else html.escape(raw);scene["blocks"].append(f"<{tag}>{body}</{tag}>");paragraph=[]
    for line in lines:
        stripped=line.strip()
        if markdown and stripped.startswith("# "):flush();chapter={"title":stripped[2:].strip(),"scenes":[]};chapters.append(chapter);scene=None;continue
        if markdown and stripped.startswith("## "):flush();ensure_chapter();ensure_scene(stripped[3:].strip());continue
        if not markdown and re.match(r"^(?:chapter|part)\s+\w+",stripped,re.I):flush();chapter={"title":stripped,"scenes":[]};chapters.append(chapter);scene=None;continue
        if SCENE_BREAK.match(stripped):flush();ensure_scene();continue
        if not stripped:flush();continue
        paragraph.append(stripped)
    flush()
    for chapter in chapters:
        for scene in chapter["scenes"]:scene["html"]="".join(scene.pop("blocks"))
    return finalize(chapters,path)

def clean_block(block):
    output=BeautifulSoup("", "html.parser").new_tag(block.name if block.name in {"p","blockquote"} else "p")
    source=BeautifulSoup(str(block),"html.parser").find(block.name)
    for child in source.descendants:
        if isinstance(child,NavigableString):
            parents=[parent.name for parent in child.parents]
            value=html.escape(str(child));
            if "strong" in parents or "b" in parents:value=f"<strong>{value}</strong>"
            if "em" in parents or "i" in parents:value=f"<em>{value}</em>"
            output.append(BeautifulSoup(value,"html.parser"))
        elif isinstance(child,Tag) and child.name=="br":output.append(BeautifulSoup("<br>","html.parser"))
    return str(output)

def inspect_html(path):
    soup=BeautifulSoup(path.read_text(encoding="utf-8-sig"),"html.parser");chapters=[];chapter=None;scene=None
    for element in soup.find_all(["h1","h2","p","blockquote","hr"]):
        if element.name=="h1":chapter={"title":element.get_text(" ",strip=True) or "Untitled Chapter","scenes":[]};chapters.append(chapter);scene=None
        elif element.name=="h2":
            if chapter is None:chapter={"title":"Imported Manuscript","scenes":[]};chapters.append(chapter)
            scene={"title":element.get_text(" ",strip=True) or f"Scene {len(chapter['scenes'])+1}","html":""};chapter["scenes"].append(scene)
        elif element.name=="hr":
            if chapter is None:chapter={"title":"Imported Manuscript","scenes":[]};chapters.append(chapter)
            scene={"title":f"Scene {len(chapter['scenes'])+1}","html":""};chapter["scenes"].append(scene)
        else:
            if chapter is None:chapter={"title":"Imported Manuscript","scenes":[]};chapters.append(chapter)
            if scene is None:scene={"title":f"Scene {len(chapter['scenes'])+1}","html":""};chapter["scenes"].append(scene)
            scene["html"]+=clean_block(element)
    return finalize(chapters,path)

def inspect_epub(path):
    chapters=[];warnings=[]
    with zipfile.ZipFile(path) as archive:
        names=[name for name in archive.namelist() if name.lower().endswith((".xhtml",".html",".htm")) and "nav" not in name.lower()]
        for name in names:
            soup=BeautifulSoup(archive.read(name),"html.parser");title=(soup.find("h1") or soup.find("title"));chapter={"title":title.get_text(" ",strip=True) if title else Path(name).stem,"scenes":[]};scene={"title":"Scene 1","html":""}
            for block in soup.find_all(["p","blockquote","hr"]):
                if block.name=="hr" and scene["html"]:chapter["scenes"].append(scene);scene={"title":f"Scene {len(chapter['scenes'])+1}","html":""}
                elif block.name!="hr":scene["html"]+=clean_block(block)
            if scene["html"]:chapter["scenes"].append(scene)
            if chapter["scenes"]:chapters.append(chapter)
    if not chapters:warnings.append("No readable XHTML manuscript content was found.")
    return finalize(chapters,path,warnings)

def inspect_document(path):
    suffix=path.suffix.lower()
    if suffix==".docx":return inspect_docx(path)
    if suffix in {".md",".markdown"}:return inspect_lines(path,True)
    if suffix==".txt":return inspect_lines(path,False)
    if suffix in {".html",".htm"}:return inspect_html(path)
    if suffix==".epub":return inspect_epub(path)
    raise ValueError(f"Unsupported import format: {suffix}")

def append_inline(paragraph,node):
    if isinstance(node,NavigableString): paragraph.add_run(str(node));return
    if not isinstance(node,Tag): return
    if node.name=="br": paragraph.add_run().add_break(WD_BREAK.LINE);return
    for child in node.children:
        before=len(paragraph.runs);append_inline(paragraph,child)
        for run in paragraph.runs[before:]:
            if node.name in {"strong","b"}:run.bold=True
            if node.name in {"em","i"}:run.italic=True
            if node.name=="u":run.underline=True

def _set_cellar_page_number(paragraph):
    """Add a PAGE field without flattening it to display text."""
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run();begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin")
    instruction=OxmlElement("w:instrText");instruction.set(qn("xml:space"),"preserve");instruction.text=" PAGE "
    separate=OxmlElement("w:fldChar");separate.set(qn("w:fldCharType"),"separate")
    end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
    for element in (begin,instruction,separate,end):run._r.append(element)

def _configure_kdp(document, payload):
    section=document.sections[0];section.page_width=Inches(6);section.page_height=Inches(9)
    section.top_margin=Inches(.75);section.bottom_margin=Inches(.75)
    section.left_margin=Inches(.875);section.right_margin=Inches(.625)
    section.header_distance=Inches(.3);section.footer_distance=Inches(.35)
    section.different_first_page_header_footer=True
    settings=document.settings._element
    if settings.find(qn("w:mirrorMargins")) is None:settings.append(OxmlElement("w:mirrorMargins"))
    normal=document.styles["Normal"];normal.font.name="Garamond";normal.font.size=Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Garamond")
    normal.paragraph_format.first_line_indent=Inches(.25);normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
    heading=document.styles["Heading 1"];heading.font.name="Garamond";heading.font.size=Pt(18);heading.font.bold=False
    heading._element.rPr.rFonts.set(qn("w:eastAsia"),"Garamond")
    heading.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;heading.paragraph_format.space_before=Inches(1.5)
    heading.paragraph_format.space_after=Pt(24);heading.paragraph_format.keep_with_next=True
    if "Lillith Scene Break" not in styles_by_name(document):
        marker=document.styles.add_style("Lillith Scene Break",WD_STYLE_TYPE.PARAGRAPH)
        marker.font.name="Garamond";marker.font.size=Pt(11);marker.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        marker.paragraph_format.space_before=Pt(10);marker.paragraph_format.space_after=Pt(10)
    footer=section.footer.paragraphs[0];_set_cellar_page_number(footer)
    document.core_properties.subject="KDP-ready 6 x 9 paperback interior"
    document.core_properties.author=payload.get("author") or ""

def _configure_hardcover(document, payload):
    section=document.sections[0];section.page_width=Inches(6);section.page_height=Inches(9)
    section.top_margin=Inches(.82);section.bottom_margin=Inches(.8)
    section.left_margin=Inches(.95);section.right_margin=Inches(.7)
    section.header_distance=Inches(.32);section.footer_distance=Inches(.36)
    section.different_first_page_header_footer=True
    settings=document.settings._element
    if settings.find(qn("w:mirrorMargins")) is None:settings.append(OxmlElement("w:mirrorMargins"))
    normal=document.styles["Normal"];normal.font.name="Garamond";normal.font.size=Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Garamond")
    normal.paragraph_format.first_line_indent=Inches(.25);normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
    heading=document.styles["Heading 1"];heading.font.name="Garamond";heading.font.size=Pt(18);heading.font.bold=False
    heading._element.rPr.rFonts.set(qn("w:eastAsia"),"Garamond")
    heading.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;heading.paragraph_format.space_before=Inches(1.55)
    heading.paragraph_format.space_after=Pt(24);heading.paragraph_format.keep_with_next=True
    if "Lillith Scene Break" not in styles_by_name(document):
        marker=document.styles.add_style("Lillith Scene Break",WD_STYLE_TYPE.PARAGRAPH)
        marker.font.name="Garamond";marker.font.size=Pt(11);marker.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        marker.paragraph_format.space_before=Pt(10);marker.paragraph_format.space_after=Pt(10)
    footer=section.footer.paragraphs[0];_set_cellar_page_number(footer)
    document.core_properties.subject="KDP-ready hardcover interior"
    document.core_properties.author=payload.get("author") or ""

def styles_by_name(document):
    return {style.name for style in document.styles}

def _add_title_page(document,payload):
    title=document.add_paragraph();title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before=Inches(2.2);title_run=title.add_run(payload.get("title") or "Manuscript")
    title_run.font.name="Garamond";title_run.font.size=Pt(24);title_run.bold=False
    if payload.get("author"):
        author=document.add_paragraph();author.alignment=WD_ALIGN_PARAGRAPH.CENTER;author.paragraph_format.space_before=Pt(24)
        author_run=author.add_run(f"by {payload['author']}");author_run.font.name="Garamond";author_run.font.size=Pt(12)
    document.add_page_break()

def export_docx(payload:dict,path:Path):
    document=Document(); styles=document.styles
    profile=payload.get("exportProfile") or "standard"
    styles["Normal"].font.name="Times New Roman";styles["Normal"].font.size=Pt(12)
    title=payload.get("title") or "Manuscript"; document.core_properties.title=title
    document.core_properties.author=payload.get("author") or ""
    if profile=="kdp-paperback":_configure_kdp(document,payload);_add_title_page(document,payload)
    elif profile=="kdp-hardcover":_configure_hardcover(document,payload);_add_title_page(document,payload)
    for chapter_index,chapter in enumerate(payload.get("chapters",[])):
        if chapter_index: document.add_page_break()
        document.add_heading(chapter.get("title") or f"Chapter {chapter_index+1}",level=1)
        for scene_index,scene in enumerate(chapter.get("scenes",[])):
            if scene_index:
                marker=document.add_paragraph("* * *",style="Lillith Scene Break" if profile in {"kdp-paperback","kdp-hardcover"} else None); marker.alignment=WD_ALIGN_PARAGRAPH.CENTER
            soup=BeautifulSoup(scene.get("html") or "<p></p>","html.parser")
            for block_index,block in enumerate(soup.find_all(["p","blockquote"],recursive=False)):
                paragraph=document.add_paragraph(style="Quote" if block.name=="blockquote" else None)
                if profile in {"kdp-paperback","kdp-hardcover"} and block.name=="p" and block_index==0:
                    paragraph.paragraph_format.first_line_indent=Inches(0)
                for child in block.children:append_inline(paragraph,child)
    path.parent.mkdir(parents=True,exist_ok=True);document.save(path)
    return {"outputPath":str(path),"chapters":len(payload.get("chapters",[])),"scenes":sum(len(ch.get("scenes",[])) for ch in payload.get("chapters",[])),"profile":profile}

def html_to_markdown(value):
    soup=BeautifulSoup(value,"html.parser");parts=[]
    for block in soup.find_all(["p","blockquote"],recursive=False):
        for tag in block.find_all(["strong","b"]):tag.replace_with(f"**{tag.get_text()}**")
        for tag in block.find_all(["em","i"]):tag.replace_with(f"*{tag.get_text()}*")
        prefix="> " if block.name=="blockquote" else "";parts.append(prefix+block.get_text(" ",strip=True))
    return "\n\n".join(parts)

def export_textual(payload,path,kind):
    output=[]
    if kind=="html":output=["<!doctype html><html><head><meta charset='utf-8'><title>",html.escape(payload.get("title") or "Manuscript"),"</title><style>body{max-width:42rem;margin:3rem auto;font:18px/1.7 Georgia,serif}h1{page-break-before:always}hr{margin:2rem;border:0;text-align:center}hr:after{content:'* * *'}blockquote{border-left:2px solid #777;padding-left:1rem}</style></head><body>"]
    for chapter in payload.get("chapters",[]):
        if kind=="html":output.append(f"<h1>{html.escape(chapter['title'])}</h1>")
        else:output.append(("# " if kind=="markdown" else "")+chapter["title"]+"\n")
        for index,scene in enumerate(chapter.get("scenes",[])):
            if index:output.append("<hr>" if kind=="html" else "\n* * *\n")
            if kind=="html":output.append(scene.get("html") or "<p></p>")
            elif kind=="markdown":output.append(html_to_markdown(scene.get("html") or ""))
            else:output.append(BeautifulSoup(scene.get("html") or "","html.parser").get_text("\n",strip=True))
    if kind=="html":output.append("</body></html>")
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text("\n\n".join(output),encoding="utf-8")

def export_epub(payload,path):
    book_id=str(uuid.uuid4());chapters=payload.get("chapters",[]);path.parent.mkdir(parents=True,exist_ok=True)
    profile=payload.get("exportProfile") or "standard"
    with zipfile.ZipFile(path,"w") as archive:
        archive.writestr("mimetype","application/epub+zip",compress_type=zipfile.ZIP_STORED);archive.writestr("META-INF/container.xml","<?xml version='1.0'?><container version='1.0' xmlns='urn:oasis:names:tc:opendocument:xmlns:container'><rootfiles><rootfile full-path='OEBPS/content.opf' media-type='application/oebps-package+xml'/></rootfiles></container>")
        manifest=[];spine=[];nav=[]
        for index,chapter in enumerate(chapters,1):
            filename=f"chapter-{index}.xhtml";body=[f"<h1>{html.escape(chapter['title'])}</h1>"]
            for scene_index,scene in enumerate(chapter.get("scenes",[])):
                if scene_index:body.append("<hr/>")
                body.append(scene.get("html") or "<p></p>")
            archive.writestr(f"OEBPS/{filename}",f"<?xml version='1.0' encoding='utf-8'?><html xmlns='http://www.w3.org/1999/xhtml'><head><title>{html.escape(chapter['title'])}</title></head><body>{''.join(body)}</body></html>");manifest.append(f"<item id='c{index}' href='{filename}' media-type='application/xhtml+xml'/>");spine.append(f"<itemref idref='c{index}'/>");nav.append(f"<li><a href='{filename}'>{html.escape(chapter['title'])}</a></li>")
        archive.writestr("OEBPS/nav.xhtml",f"<?xml version='1.0' encoding='utf-8'?><html xmlns='http://www.w3.org/1999/xhtml' xmlns:epub='http://www.idpf.org/2007/ops'><head><title>Contents</title></head><body><nav epub:type='toc'><ol>{''.join(nav)}</ol></nav></body></html>");manifest.append("<item id='nav' href='nav.xhtml' media-type='application/xhtml+xml' properties='nav'/>")
        archive.writestr("OEBPS/content.opf",f"<?xml version='1.0' encoding='utf-8'?><package xmlns='http://www.idpf.org/2007/opf' version='3.0' unique-identifier='bookid'><metadata xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:identifier id='bookid'>urn:uuid:{book_id}</dc:identifier><dc:title>{html.escape(payload.get('title') or 'Manuscript')}</dc:title><dc:language>en</dc:language><meta name='lillith:export-profile' content='{html.escape(profile)}'/></metadata><manifest>{''.join(manifest)}</manifest><spine>{''.join(spine)}</spine></package>")

def export_pdf(payload,path):
    lines=[payload.get("title") or "Manuscript",""]
    for chapter in payload.get("chapters",[]):
        lines.extend([chapter["title"],""])
        for index,scene in enumerate(chapter.get("scenes",[])):
            if index:lines.extend(["* * *",""])
            plain=BeautifulSoup(scene.get("html") or "","html.parser").get_text(" ",strip=True)
            lines.extend(textwrap.wrap(plain,width=82) or [""]);lines.append("")
    pages=[lines[index:index+48] for index in range(0,len(lines),48)] or [[]];objects=[b"",b""]
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Times-Roman >>")
    for page_index,page in enumerate(pages):
        content=["BT /F1 11 Tf 54 760 Td 14 TL"]
        for line in page:
            safe=line.encode("latin-1","replace").decode("latin-1").replace("\\","\\\\").replace("(","\\(").replace(")","\\)");content.append(f"({safe}) Tj T*")
        content.append("ET");stream="\n".join(content).encode("latin-1");page_obj=4+page_index*2;content_obj=page_obj+1;objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {content_obj} 0 R >>".encode());objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode()+stream+b"\nendstream")
    kids=" ".join(f"{4+i*2} 0 R" for i in range(len(pages)));objects[0]=b"<< /Type /Catalog /Pages 2 0 R >>";objects[1]=f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode();data=bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n");offsets=[0]
    for number,obj in enumerate(objects,1):offsets.append(len(data));data.extend(f"{number} 0 obj\n".encode()+obj+b"\nendobj\n")
    xref=len(data);data.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode());
    for offset in offsets[1:]:data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode());path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(data)

def export_document(payload,path):
    suffix=path.suffix.lower()
    profile=payload.get("exportProfile") or "standard"
    if suffix==".docx":result=export_docx(payload,path)
    elif suffix in {".md",".markdown"}:export_textual(payload,path,"markdown");result={}
    elif suffix==".txt":export_textual(payload,path,"text");result={}
    elif suffix in {".html",".htm"}:export_textual(payload,path,"html");result={}
    elif suffix==".epub":export_epub(payload,path);result={}
    elif suffix==".pdf":export_pdf(payload,path);result={}
    else:raise ValueError(f"Unsupported export format: {suffix}")
    return {"outputPath":str(path),"chapters":len(payload.get("chapters",[])),"scenes":sum(len(ch.get("scenes",[])) for ch in payload.get("chapters",[])),"profile":profile,**result}

def main():
    parser=argparse.ArgumentParser();parser.add_argument("operation",choices=["inspect","export"]);parser.add_argument("--input");parser.add_argument("--output");args=parser.parse_args()
    try:
        if args.operation=="inspect": result=inspect_document(Path(args.input).resolve())
        else: result=export_document(json.load(sys.stdin),Path(args.output).resolve())
        print(json.dumps({"ok":True,"result":result},ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok":False,"error":f"{type(exc).__name__}: {exc}"}),file=sys.stdout);sys.exit(1)
if __name__=="__main__":main()
