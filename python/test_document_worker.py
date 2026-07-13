import tempfile, unittest, zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from document_worker import inspect_docx, export_docx, inspect_document, export_document

class DocumentWorkerTests(unittest.TestCase):
    def test_docx_round_trip_preserves_structure_and_emphasis(self):
        with tempfile.TemporaryDirectory() as folder:
            source=Path(folder)/"source.docx";doc=Document();doc.add_heading("Chapter One",1);p=doc.add_paragraph();p.add_run("Dark water").bold=True;p.add_run(" moved.").italic=True;doc.add_paragraph("* * *");doc.add_paragraph("Second scene");doc.save(source)
            preview=inspect_docx(source);self.assertEqual(preview["counts"]["chapters"],1);self.assertEqual(preview["counts"]["scenes"],2);self.assertIn("<strong>Dark water</strong>",preview["chapters"][0]["scenes"][0]["html"])
            output=Path(folder)/"output.docx";export_docx({"title":"Test","chapters":preview["chapters"]},output);loaded=Document(output);self.assertEqual(loaded.paragraphs[0].text,"Chapter One");self.assertTrue(any(run.bold for p in loaded.paragraphs for run in p.runs if "Dark water" in run.text))
    def test_markdown_html_epub_and_pdf_outputs(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/"story.md";source.write_text("# Chapter One\n\n## Low Water\n\n**Dark water** moved.\n\n* * *\n\n> The bell sounded.",encoding="utf-8");preview=inspect_document(source);self.assertEqual(preview["counts"]["scenes"],2);self.assertIn("<strong>Dark water</strong>",preview["chapters"][0]["scenes"][0]["html"])
            payload={"title":"Test","chapters":preview["chapters"]}
            for extension in ["md","txt","html","epub","pdf"]:
                output=root/f"output.{extension}";result=export_document(payload,output);self.assertTrue(output.is_file());self.assertGreater(output.stat().st_size,20);self.assertEqual(result["chapters"],1)
            self.assertTrue((root/"output.pdf").read_bytes().startswith(b"%PDF-1.4"))
            with zipfile.ZipFile(root/"output.epub") as archive:self.assertEqual(archive.read("mimetype"),b"application/epub+zip");self.assertIn("OEBPS/content.opf",archive.namelist())
    def test_kdp_docx_uses_print_layout_and_keeps_manuscript_formatting(self):
        with tempfile.TemporaryDirectory() as folder:
            output=Path(folder)/"paperback.docx"
            payload={"title":"The Wake","author":"M. Vale","exportProfile":"kdp-paperback","chapters":[{"title":"Chapter One","scenes":[{"title":"Arrival","html":"<p><strong>Dark water</strong> moved.</p><p>The bell answered.</p>"},{"title":"Below","html":"<p><em>Again.</em></p>"}]}]}
            result=export_document(payload,output);loaded=Document(output);section=loaded.sections[0]
            self.assertEqual(result["profile"],"kdp-paperback")
            self.assertAlmostEqual(section.page_width.inches,6,places=2);self.assertAlmostEqual(section.page_height.inches,9,places=2)
            self.assertAlmostEqual(section.left_margin.inches,.875,places=2);self.assertAlmostEqual(section.right_margin.inches,.625,places=2)
            self.assertIsNotNone(loaded.settings._element.find(qn("w:mirrorMargins")))
            self.assertEqual(loaded.core_properties.author,"M. Vale");self.assertIn("KDP-ready",loaded.core_properties.subject)
            self.assertEqual(loaded.paragraphs[0].text,"The Wake");self.assertTrue(any(p.text=="* * *" for p in loaded.paragraphs))
            self.assertTrue(any(run.bold for p in loaded.paragraphs for run in p.runs if "Dark water" in run.text))
            first_body=next(p for p in loaded.paragraphs if p.text=="Dark water moved.")
            second_body=next(p for p in loaded.paragraphs if p.text=="The bell answered.")
            self.assertEqual(first_body.paragraph_format.first_line_indent,Inches(0))
            self.assertEqual(second_body.style.paragraph_format.first_line_indent,Inches(.25))
    def test_publication_profiles_cover_kindle_and_hardcover_routes(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            payload={"title":"The Wake","author":"M. Vale","exportProfile":"kindle","chapters":[{"title":"Chapter One","scenes":[{"title":"Arrival","html":"<p>Dark water moved.</p>"}]}]}
            kindle=root/"kindle.epub"; hardback=root/"hardcover.docx"
            kindle_result=export_document(payload,kindle); hardcover_result=export_document({**payload,"exportProfile":"kdp-hardcover"},hardback)
            self.assertEqual(kindle_result["profile"],"kindle")
            self.assertEqual(hardcover_result["profile"],"kdp-hardcover")
            self.assertTrue(kindle.is_file()); self.assertTrue(hardback.is_file())
            with zipfile.ZipFile(kindle) as archive:self.assertIn("OEBPS/content.opf",archive.namelist())
            loaded=Document(hardback); self.assertIn("KDP-ready hardcover",loaded.core_properties.subject)
if __name__=="__main__":unittest.main()
